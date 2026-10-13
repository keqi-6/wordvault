"""真题语料 -> 词频、真题例句、首期点亮词。

数据来源（实测）：
  data/raw/CET6-Resources-full/   完整真题库（253 MB tar 解压后 311 MB，344 个文件）
      - 62 个 .docx 可以直接抽出正文（每年真题一份，含阅读原文、听力文字稿、翻译、答案）
      - 217 个 PDF 用 pypdf 试抽；扫描版抽不出文字的会被统计出来
      - 34 个 .doc 是旧二进制格式，Python 读不了，单独报告
  data/raw/CET6-Resources/        之前的稀疏检出（只有几个 md，作为补充）

输出: data/build/exam.json
"""
from __future__ import annotations

import re
import sys
from collections import defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import (  # noqa: E402
    BUILD, DICT_JSON, EXAM_DIR, LIT_LIMIT, RAW, ensure_dirs, log, norm_word, read_json, write_json,
)

FULL_DIR = RAW / "CET6-Resources-full" / "CET6-Resources-master"
TEXT_EXT = {".md", ".txt", ".csv", ".tsv", ".rtf"}
DOCX_EXT = {".docx"}
DOC_EXT = {".doc"}
PDF_EXT = {".pdf"}
IMG_EXT = {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp"}
TOKEN = re.compile(r"[A-Za-z][A-Za-z'\-]*")
SENT_SPLIT = re.compile(r"(?<=[.!?])\s+(?=[A-Z\"'(])")
CJK = re.compile(r"[\u3000-\u9fff\uff00-\uffef]")
OPTION_PREFIX = re.compile(r"^\s*(?:[A-D][).\uff09]|\d{1,2}\s*[.\uff0e\u3001])\s*")
SKIP_DIRS = {".git", "node_modules", "__pycache__"}
VOCAB_HINT_FILES = ("可用词汇整理.md", "可用词汇整理.txt")

# 真题里每套卷子都重复一遍的指令性文字。不过滤掉的话，
# 词频榜会被 corresponding / sheet / direction 这类"考题用语"霸占，
# 而它们并不是要背的词汇。
BOILERPLATE = re.compile(
    r"(directions?\s*:|answer sheet|corresponding letter|you will hear|"
    r"best answer|single line through|record your answer|questions?\s+\d+\s+to\s+\d+|"
    r"^\s*section\s+[ABC]\b|^\s*part\s+[IVX]+\b|for this part|"
    r"write an essay|at least \d+ words|no more than \d+ words|"
    r"suppose you|you should write|write your answer|complete the sentences|"
    r"reading comprehension|listening comprehension|"
    r"\(\s*\d+\s*minutes?\s*\)|marked A\)|choices marked)",
    re.IGNORECASE)


def inflection_map(words: list[dict]) -> dict[str, str]:
    out: dict[str, str] = {}
    for w in words:
        ex = w.get("exchange") or ""
        base = w["lemma"]
        if not ex:
            continue
        for part in ex.split("/"):
            if ":" not in part:
                continue
            k, v = part.split(":", 1)
            if k == "0":
                if v:
                    out.setdefault(norm_word(v), base)
                out.setdefault(base, base)
            elif k in {"s", "p", "d", "i", "3", "r", "t"}:
                for form in v.split(","):
                    form = norm_word(form)
                    if form and form != base:
                        out.setdefault(form, base)
    return out


def read_docx(path: Path) -> str:
    try:
        import docx  # python-docx
    except ImportError:
        return ""
    try:
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append("\t".join(c.text for c in row.cells))
        return "\n".join(parts)
    except Exception:
        return ""


def read_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        return ""
    try:
        reader = PdfReader(str(path))
        max_pages = min(len(reader.pages), 40)
        return "\n".join((reader.pages[i].extract_text() or "") for i in range(max_pages))
    except Exception:
        return ""


def ocr_image(path: Path) -> str:
    try:
        import pytesseract  # type: ignore
        from PIL import Image
    except ImportError:
        return ""
    try:
        return pytesseract.image_to_string(Image.open(path), lang="eng")
    except Exception:
        return ""


def collect_text() -> tuple[str, dict, list[Path]]:
    chunks: list[str] = []
    hint_files: list[Path] = []
    stats = defaultdict(int)
    roots = [d for d in (FULL_DIR, EXAM_DIR) if d.exists()]
    if not roots:
        log("  未找到真题目录")
        return "", dict(stats), []
    files = [p for r in roots for p in r.rglob("*") if p.is_file() and not any(s in p.parts for s in SKIP_DIRS)]
    log(f"  真题目录文件数: {len(files)}")
    for p in files:
        ext = p.suffix.lower()
        if not ext:
            continue
        if ext in DOCX_EXT:
            t = read_docx(p)
            stats["docx"] += 1
            if t.strip():
                stats["docx_ok"] += 1
                chunks.append(t)
        elif ext in TEXT_EXT:
            try:
                chunks.append(p.read_text(encoding="utf-8", errors="replace"))
                stats["text"] += 1
                if p.name in VOCAB_HINT_FILES:
                    hint_files.append(p)
            except Exception:
                stats["text_fail"] += 1
        elif ext in PDF_EXT:
            t = read_pdf(p)
            stats["pdf"] += 1
            if len(t.strip()) > 200:
                stats["pdf_text"] += 1
                chunks.append(t)
            else:
                stats["pdf_scan"] += 1
        elif ext in DOC_EXT:
            stats["doc_legacy"] += 1
        elif ext in IMG_EXT:
            t = ocr_image(p)
            stats["img"] += 1
            if t.strip():
                stats["img_ocr"] += 1
                chunks.append(t)
    log("  解析统计: " + str(dict(stats)))
    return "\n".join(chunks), dict(stats), hint_files


def main() -> None:
    ensure_dirs()
    doc = read_json(DICT_JSON)
    if not doc:
        raise SystemExit("请先运行 parse_dict.py")
    words = doc["words"]
    vocab = {w["lemma"] for w in words}
    cet6_set = {w["lemma"] for w in words if w.get("in_cet6")}
    infl = inflection_map(words)
    by_lemma = {w["lemma"]: w for w in words}
    log("词库 " + str(len(vocab)) + " 词（六级 " + str(len(cet6_set)) + "），变形映射 " + str(len(infl)) + " 条")

    text, stats, hint_files = collect_text()
    log("  文本长度: " + str(len(text)) + " 字符")

    freq: dict[str, int] = defaultdict(int)
    sentences: dict[str, list[dict]] = defaultdict(list)
    if text.strip():
        dropped = 0
        for sent in SENT_SPLIT.split(text):
            s = " ".join(sent.split())
            if not (40 <= len(s) <= 320):
                continue
            if BOILERPLATE.search(s) or s.count("A)") + s.count("B)") >= 2:
                dropped += 1
                continue
            # 听力文字稿后面常跟着中文标注（【预测】…）或下一题的中文解析，
            # 直接在第一个中日韩字符处截断，英文部分不够长就丢弃。
            cjk = CJK.search(s)
            if cjk:
                s = s[:cjk.start()].rstrip()
                if len(s) < 40:
                    dropped += 1
                    continue
            s = OPTION_PREFIX.sub("", s).strip()
            if len(s) < 40:
                dropped += 1
                continue
            hit: set[str] = set()
            for tok in TOKEN.findall(s):
                t = norm_word(tok)
                base = t if t in vocab else infl.get(t)
                if base and base in vocab:
                    freq[base] += 1
                    hit.add(base)
            for b in list(hit)[:8]:
                if len(sentences[b]) < 5:
                    sentences[b].append({"en": s, "source": "cet-exam"})

    log("  过滤掉的指令性句子: " + str(dropped))

    exam_vocab: set[str] = set()
    for p in hint_files:
        try:
            raw = p.read_text(encoding="utf-8", errors="replace")
        except Exception:
            continue
        for tok in TOKEN.findall(raw):
            t = norm_word(tok)
            base = t if t in vocab else infl.get(t)
            if base and base in cet6_set:
                exam_vocab.add(base)

    def rank_key(lemma: str) -> tuple:
        e = by_lemma.get(lemma, {})
        return (-freq.get(lemma, 0), e.get("frq") or 99999, lemma)

    # 中考/高考级别的词（ECDICT 打了 zk/gk 标签）不放进「要背的 500 词」，
    # 否则真题里最常见的 what / can / way 会把名额吃光。
    basic = {w["lemma"] for w in words
             if any(t in ("zk", "gk") for t in (w.get("tag") or "").split())}
    ranked = [w for w in sorted(cet6_set, key=rank_key) if w not in basic]
    lit = ranked[:LIT_LIMIT]
    log("  中考/高考级基础词（不入选）: " + str(len(basic)))
    log("  有真题词频的六级词: " + str(sum(1 for w in cet6_set if freq.get(w))))
    log("  可用词汇整理命中: " + str(len(exam_vocab)))
    log("  首期点亮: " + str(len(lit)) + " 词（按真题出现次数排）")

    out = {
        "stats": stats,
        "text_chars": len(text),
        "freq": dict(freq),
        "sentences": dict(sentences),
        "exam_vocab": sorted(exam_vocab),
        "lit": lit,
        "note": "词频来自完整真题库（docx 正文优先，PDF 抽到文字的一并计入）；lit 按真题出现次数排序",
    }
    write_json(BUILD / "exam.json", out)
    log("写出 " + str(BUILD / "exam.json"))
    top = [(w, freq.get(w, 0)) for w in lit[:25]]
    log("真题高频前 25: " + ", ".join(f"{w}({c})" for w, c in top))


if __name__ == "__main__":
    main()
